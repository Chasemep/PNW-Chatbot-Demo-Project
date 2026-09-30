"""DOCX parsing into ordered, structure-aware blocks."""

from pathlib import Path
from subprocess import CalledProcessError, TimeoutExpired, run
from tempfile import TemporaryDirectory

from docx import Document
from docx.document import Document as WordDocument
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table
from docx.text.paragraph import Paragraph

from app.models.base import ContentKind
from app.services.ingestion.parser import ParsedDocument, StructuralBlock


def parse_docx(path: Path) -> ParsedDocument:
    """Parse DOCX headings, content, tables, lists, sidebars, and callouts."""

    if path.suffix.lower() == ".doc":
        return _parse_legacy_doc(path)
    if path.suffix.lower() != ".docx":
        raise ValueError(f"unsupported Word document format: {path.suffix}")
    return _parse_document(Document(str(path)))


def _parse_legacy_doc(path: Path) -> ParsedDocument:
    """Convert a legacy DOC file before applying the DOCX structural parser."""

    with TemporaryDirectory() as output_directory:
        try:
            run(
                [
                    "soffice",
                    "--headless",
                    "--convert-to",
                    "docx",
                    "--outdir",
                    output_directory,
                    str(path),
                ],
                check=True,
                capture_output=True,
                text=True,
                timeout=30,
            )
        except (CalledProcessError, OSError, TimeoutExpired) as error:
            raise ValueError("could not convert legacy DOC source") from error

        converted_path = Path(output_directory) / f"{path.stem}.docx"
        if not converted_path.is_file():
            raise ValueError("legacy DOC conversion did not produce a DOCX file")
        return _parse_document(Document(str(converted_path)))


def _parse_document(document: WordDocument) -> ParsedDocument:
    """Parse an opened Word document into ordered structural blocks."""

    headings: list[str] = []
    blocks: list[StructuralBlock] = []
    list_items: list[str] = []

    def append_block(kind: ContentKind, text: str) -> None:
        normalized = " ".join(text.split())
        if not normalized:
            return
        blocks.append(
            StructuralBlock(
                content_kind=kind,
                content_text=normalized,
                structural_path=" > ".join(headings) or "Document",
                location_label=f"DOCX element {len(blocks) + 1}",
                ordinal=len(blocks),
            )
        )

    def flush_list() -> None:
        if list_items:
            append_block(ContentKind.LIST, "\n".join(list_items))
            list_items.clear()

    for element in document.element.body.iterchildren():
        if isinstance(element, CT_P):
            paragraph = Paragraph(element, document)
            text = paragraph.text.strip()
            style_name = paragraph.style.name if paragraph.style is not None else ""
            if not text:
                continue
            if style_name.startswith("Heading "):
                flush_list()
                level = int(style_name.removeprefix("Heading "))
                del headings[level - 1 :]
                headings.append(text)
                append_block(ContentKind.HEADING, text)
            elif style_name.startswith("List "):
                list_items.append(text)
            else:
                flush_list()
                if style_name == "Sidebar":
                    append_block(ContentKind.SIDEBAR, text)
                elif style_name == "Callout":
                    append_block(ContentKind.CALLOUT, text)
                else:
                    append_block(ContentKind.PARAGRAPH, text)
        elif isinstance(element, CT_Tbl):
            flush_list()
            table = Table(element, document)
            rows = [
                " | ".join(cell.text.strip() for cell in row.cells)
                for row in table.rows
            ]
            append_block(ContentKind.TABLE, "\n".join(rows))

    flush_list()
    return ParsedDocument(blocks=blocks)
